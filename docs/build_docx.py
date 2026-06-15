#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""產生 Word .docx 版期末報告，內容與 PDF 一致。"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

DOCS = "/home/user/final_project/docs"
SHOTS = os.path.join(DOCS, "screenshots")
OUT = "/home/user/final_project/report/4B4G0039.docx"
os.makedirs(os.path.dirname(OUT), exist_ok=True)

CJK = "Microsoft JhengHei"   # 微軟正黑體（Windows 內建，支援繁體中文）
BLUE = RGBColor(0x2b, 0x4c, 0x7e)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# 預設字型（含東亞字型）
style = doc.styles["Normal"]
style.font.name = CJK
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

def set_cjk(run):
    run.font.name = CJK
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), CJK)

def para(text="", size=11, bold=False, color=None, align=None, italic=False,
         space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        set_cjk(r)
    return p

def heading(text, size=15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = BLUE
    set_cjk(r)
    # 底線（段落下框線）
    pPr = p._p.get_or_add_pPr()
    pbdr = pPr.makeelement(qn("w:pBdr"), {})
    bottom = pbdr.makeelement(qn("w:bottom"),
        {qn("w:val"): "single", qn("w:sz"): "12",
         qn("w:space"): "2", qn("w:color"): "4a6fa5"})
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p

def subheading(text, size=12.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = BLUE
    set_cjk(r)
    return p

def code_block(lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    shade_para(p, "F2F2F2")
    for i, ln in enumerate(lines):
        r = p.add_run(("" if i == 0 else "\n") + ln)
        r.font.name = "Consolas"
        r.font.size = Pt(10)
        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK)
    return p

def shade_para(p, hexcolor):
    pPr = p._p.get_or_add_pPr()
    shd = pPr.makeelement(qn("w:shd"),
        {qn("w:val"): "clear", qn("w:fill"): hexcolor})
    pPr.append(shd)

def table(rows, widths_in, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Inches(widths_in[ci])
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            r = cell.paragraphs[0].add_run(str(val))
            r.font.size = Pt(10.5)
            r.bold = (header and ri == 0)
            set_cjk(r)
            if header and ri == 0:
                shade_para(cell.paragraphs[0], "DDE6F2")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

def pic(path, width_in):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width_in))
    return p

C = WD_ALIGN_PARAGRAPH.CENTER

# ===================== 封面 =====================
for _ in range(3):
    para("")
para("物件導向程式設計　期末專題報告", size=14, color=GREY, align=C)
para("商店庫存管理系統", size=30, bold=True, color=RGBColor(0x1f,0x3a,0x63),
     align=C, space_after=2)
para("Store Inventory Management System", size=13, color=RGBColor(0x77,0x77,0x77),
     align=C, space_after=24)
para("")

# 資訊框（1x1 表格）
info = ["指導老師：陳定宏", "學　號：4B4G0039",
        "姓　名：江金桔", "製作日期：2026 / 06 / 15"]
bt = doc.add_table(rows=1, cols=1)
bt.style = "Table Grid"
bt.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = bt.cell(0, 0)
cell.width = Inches(3.6)
first = True
for line in info:
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    first = False
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(line)
    r.font.size = Pt(13)
    set_cjk(r)

for _ in range(3):
    para("")
para("GitHub Repo", size=12, color=GREY, align=C, space_after=2)
para("https://github.com/jamiechu218/final_project", size=12,
     color=RGBColor(0x1f,0x3a,0x63), align=C)
doc.add_page_break()

# ===================== 一、專題簡介 =====================
heading("一、專題簡介")
para("本專題以 C++ 物件導向設計一套於終端機 (Terminal) 操作的商店庫存管理系統。"
     "商店中的商品分為「食品、電子產品、服飾」三大類，使用者可透過文字選單進行新增、"
     "查詢、進貨、銷售、修改、刪除與統計，所有資料以檔案永久保存，下次開啟程式時自動載入。")
subheading("對應題目六大要求")
table([
    ["要求", "本專題實作方式"],
    ["1. 呈現 C++ 類別繼承", "抽象基底類別 Product 衍生 Food、Electronics、Clothing，並以虛擬函式達成多型"],
    ["2. 讀檔與寫檔", "Inventory::saveToFile() 寫檔、loadFromFile() 讀檔；啟動自動載入、離開自動存檔"],
    ["3. 使用 STL 類別庫", "vector、map、unique_ptr、string、algorithm、sstream"],
    ["4. 終端機 UI 介面", "main.cpp 提供 12 項功能的文字選單與輸入錯誤處理"],
    ["5. 規格與開發流程", "docs/SPECIFICATION.md、docs/DEVELOPMENT.md"],
    ["6. 完整說明文件", "README、規格書、開發流程與本份報告"],
], [2.0, 4.5])

# ===================== 二、類別繼承架構 =====================
heading("二、類別繼承架構（重點）")
para("系統核心是一個三層的繼承架構。Product 為抽象基底類別，定義所有商品共用的屬性與"
     "三個純虛擬函式；三個衍生類別各自擁有專屬屬性並覆寫虛擬函式：")
code_block([
    "        Product  (抽象基底類別)",
    "        |- id_  name_  price_  quantity_",
    "        |- category()      = 0   <- 純虛擬",
    "        |- displayDetail() = 0   <- 純虛擬 (多型)",
    "        |- serialize()     = 0   <- 純虛擬",
    "                 |  public 繼承",
    "      +----------+-----------+",
    "    Food     Electronics   Clothing",
    "  (有效期限)   (保固月數)    (尺寸)",
])
para("Inventory 以 std::vector<std::unique_ptr<Product>> 持有所有商品，僅透過基底指標"
     "呼叫 displayDetail()，即由實際型別決定輸出格式，這就是多型 (polymorphism)。")

# ===================== 三、功能說明 =====================
heading("三、功能說明")
table([
    ["編號", "功能", "說明"],
    ["1", "新增商品", "選擇類別後輸入名稱、單價、數量與專屬欄位，自動配發編號"],
    ["2", "顯示所有商品", "依實際類別以不同格式列出（多型展示）"],
    ["3", "依類別顯示", "只列出食品 / 電子 / 服飾其中一類"],
    ["4", "搜尋商品", "可依名稱關鍵字或依商品編號搜尋"],
    ["5", "進貨", "對指定編號增加庫存"],
    ["6", "銷售", "減少庫存；庫存不足則拒絕，避免負庫存"],
    ["7", "修改價格", "變更指定商品單價"],
    ["8", "刪除商品", "以 std::remove_if 移除指定商品"],
    ["9", "統計報表", "以 std::map 統計各類別種類數與庫存總值"],
    ["10", "儲存至檔案", "將庫存寫入 data/inventory.txt"],
    ["11", "從檔案讀取", "重新由檔案載入庫存"],
    ["0", "離開系統", "自動存檔後結束"],
], [0.7, 1.4, 4.4])
doc.add_page_break()

# ===================== 四、程式執行畫面 =====================
heading("四、程式執行畫面與說明")
subheading("畫面 1：啟動載入資料 + 主選單（刷新式單頁 UI）")
para("程式啟動時自動由 data/inventory.txt 讀入 6 筆商品，按 Enter 後進入主選單。"
     "採刷新式單頁介面，每次操作後會清空畫面、重畫乾淨的選單。",
     size=10.5, color=GREY)
pic(os.path.join(SHOTS, "01_startup_menu.png"), 3.4)
subheading("畫面 2：銷售 — 列出清單用「項次」選取")
para("進貨 / 銷售 / 改價 / 刪除時會先列出商品清單，使用者只需輸入「項次」(1,2,3…) 即可選取，"
     "不必記憶商品編號；清單同時以不同格式顯示三種商品，即為多型。完成後回報最新庫存。",
     size=10.5, color=GREY)
pic(os.path.join(SHOTS, "02_sell_select.png"), 6.2)
doc.add_page_break()

subheading("畫面 3：搜尋 — 可依名稱或依商品編號")
para("搜尋商品時可選擇「依名稱關鍵字」或「依商品編號」兩種方式，圖中示範以編號 1004 找到行動電源。",
     size=10.5, color=GREY)
pic(os.path.join(SHOTS, "03_search.png"), 5.6)
subheading("畫面 4：防呆與取消機制")
para("必填欄位若只按 Enter（空白）會被擋下並要求重新輸入（防呆）；"
     "任何輸入步驟輸入 q 可隨時取消目前動作並返回主選單。",
     size=10.5, color=GREY)
pic(os.path.join(SHOTS, "04_foolproof_cancel.png"), 5.4)

# ===================== 五、編譯與執行 =====================
heading("五、編譯與執行方式")
code_block([
    "$ make          # 以 g++ (C++14) 編譯, -Wall -Wextra 無警告",
    "$ ./inventory   # 執行 (或 make run)",
    "# Windows 可直接雙擊 build.bat 一鍵編譯並啟動",
])

# ===================== 六、結語 =====================
heading("六、結語")
para("本專題完整實踐了 C++ 的類別繼承與多型、STL 容器、檔案讀寫以及終端機 UI。"
     "透過抽象基底類別與虛擬函式的設計，未來若要新增商品類別（例如「書籍」），"
     "只需新增一個衍生類別並實作三個虛擬函式即可，不必改動既有程式，"
     "展現了物件導向「易擴充、易維護」的優點。")
para("完整原始碼：https://github.com/jamiechu218/final_project",
     bold=True, color=BLUE, align=C)

doc.save(OUT)
print("DOCX saved:", OUT)
